#!/usr/bin/env python3
"""Build the Final Review (Review-III) deck from the Review-2 deck template.

usage: make_final_deck.py [BASE.pptx] [OUT.pptx]
Numbers come from results/final_numbers.md, results/loop/campaign/summary.csv,
results/figdata.json and results/comp/mc_v2_analysis.txt (see FACTS3 rules).
"""
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import make_deck_charts as H  # noqa: E402
from make_deck_charts import (DARK, DIA, FIG, GREEN, GREY, NAVY, ORANGE, PURPLE, RED, ROOT,  # noqa: E402
                              TEAL, W, X0, Y0, CC, ORDER, bar, error_bars, legend, new_slide,
                              picture, pts, stat, text, xy)
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Pt  # noqa: E402

REPO = ROOT.parent
BASE = REPO / "Design_of_a_Frequency_Synthesiser_Review-2.pptx"
OUT = REPO / "Design_of_a_Frequency_Synthesiser_Final_Review.pptx"
D = H.D

# ---------- data ----------
def md_table(section):
    txt = (ROOT / "results/final_numbers.md").read_text()
    blk = txt.split(section, 1)[1].split("\n\n", 2)[1]
    rows = [[c.strip() for c in l.strip("|").split("|")] for l in blk.splitlines()]
    return [dict(zip(rows[0], r)) for r in rows[2:]]

COMP = md_table("## Compensation on 11 corners")
LOOP = {r["case"]: r for r in md_table("## Closed loop")}
LABEL = {"nsps": "N slow / P slow (ss)", "nspt": "N slow / P typ", "nspf": "N slow / P fast",
         "ntps": "N typ / P slow", "ntpt": "N typ / P typ (tt)", "ntpf": "N typ / P fast",
         "nfps": "N fast / P slow", "nfpt": "N fast / P typ", "nfpf": "N fast / P fast (ff)",
         "pdk_sf": "PDK sf (N fast / P slow)", "pdk_fs": "PDK fs (N slow / P fast)"}
SHORT = {"nsps": "NsPs", "nspt": "NsPt", "nspf": "NsPf", "ntps": "NtPs", "ntpt": "NtPt", "ntpf": "NtPf",
         "nfps": "NfPs", "nfpt": "NfPt", "nfpf": "NfPf", "pdk_sf": "PDK sf", "pdk_fs": "PDK fs"}
KEYS = list(LABEL)
assert [r["corner"] for r in COMP] == KEYS
CRN = {r["corner"]: r for r in COMP}

GREYTXT = RGBColor(0x77, 0x77, 0x77)
LIGHT = RGBColor(0xF2, 0xF4, 0xFB)


# ---------- helpers ----------
def takeaway(s, msg, y=5800000):
    text(s, X0, y, W, 420000, [("Takeaway: " + msg, True, NAVY)], size=15)


def table(s, x, y, w, rows, colw, size=11, rowh=300000, bad=(), good=(), left=1):
    gf = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, rowh * len(rows))
    t = gf.table
    tot = sum(colw)
    for j, c in enumerate(colw):
        t.columns[j].width = int(w * c / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = rowh
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.margin_top = c.margin_bottom = 20000
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xFB, 0xE9, 0xE7) if i in bad else
                                                         RGBColor(0xE6, 0xF4, 0xEC) if i in good else
                                                         LIGHT if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j < left else PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(v)
            r.font.size, r.font.name, r.font.bold = Pt(size), "Calibri", i == 0
            r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else DARK
    return t


def fig_slide(prs, logo, title, png, stats, take, side=None):
    """Figure + stat boxes + takeaway. stats: [(big, small, colour)]."""
    s = new_slide(prs, title, logo)
    from PIL import Image
    w, h = Image.open(png).size
    if w / h > 1.8 and stats:  # wide figure: stats in a row underneath
        picture(s, png, X0, Y0 - 50000, W, 3050000)
        n = len(stats); gap = 150000; bw = (W - gap * (n - 1)) // n
        for i, (b, sm, c) in enumerate(stats):
            stat(s, X0 + i * (bw + gap), 4450000, bw, 1200000, b, sm, c, size=22)
    else:  # tall figure: stats stacked on the right
        picture(s, png, X0, Y0 - 50000, 6900000, 4350000)
        bh = 1050000
        for i, (b, sm, c) in enumerate(stats):
            stat(s, 7800000, Y0 + i * (bh + 130000), 3720000, bh, b, sm, c, size=22)
    if take:
        takeaway(s, take)
    return s


def txt_slide(prs, logo, title, items, size=14, take=None, w=W):
    s = new_slide(prs, title, logo)
    text(s, X0, Y0, w, 4400000, items, size=size)
    if take:
        takeaway(s, take)
    return s


def hd(t):
    return (t, True, NAVY)


# ---------- build ----------
def build(prs, logo):
    S = []
    add = S.append
    P = prs

    # 2 problem & objectives
    s = new_slide(P, "Problem & Objectives", logo)
    text(s, X0, Y0, 6900000, 4400000, [
        hd("Problem"),
        "A ring-VCO drifts with temperature: threshold voltage, mobility and capacitance all shift, so the same control voltage gives a different frequency hot vs cold. Process corners change the size of the drift.",
        hd("Objectives"),
        "Integer-N charge-pump PLL in sky130B, open-source and headless (ngspice 47): REF 28.275 MHz -> PFD -> charge pump -> loop filter -> 5-stage current-starved ring VCO -> /8.",
        "Measure the uncompensated drift on every process corner, design a temperature-compensation, and validate it with Monte Carlo and hypothesis tests.",
        "Close the loop and check lock time across corners and temperature."], size=13)
    for i, (b, sm, c) in enumerate([("226.2 MHz", "target at V_ctrl = 0.9 V (tt, 27 C) = 8 x 28.275 MHz", NAVY),
                                    ("+/-0.3 %", "frequency-drift target over -40 to 85 C", TEAL),
                                    ("< 2 us", "lock-time target", ORANGE)]):
        stat(s, 7800000, Y0 + i * 1180000, 3720000, 1050000, b, sm, c, size=24)
    takeaway(s, "we need a drift figure measured on the real topology and process, then a correction that is proven, not assumed.")
    add(s)

    # 5 novelty
    s = new_slide(P, "Novelty, Patent and Research Gap", logo)
    text(s, X0, Y0, 6900000, 4400000, [
        hd("What is NOT novel"),
        "Generic quadratic (2nd-order) temperature compensation of an oscillator is patented: US 9,748,962 B2 (Ortner and Punzenberger, Infineon, 2017) [12]. We do not claim it.",
        "Current-starved ring VCOs and the charge-pump PLL topology are textbook.",
        hd("What this project adds"),
        "A closed-loop, per-corner, statistically validated study of a compensated ring-VCO PLL on an open PDK (sky130B), with the failures and corrections documented.",
        hd("Research gap (from 15 IEEE-level references)"),
        "No reported work combines a compensated ring VCO inside a closed loop, characterised across process corners with Monte Carlo rigour."], size=13)
    stat(s, 7800000, Y0, 3720000, 1250000, "15 refs", "read; closest: Souliotis 2022 [1], Stancu 2023 [2], Geng 2018 [3]", NAVY, size=24)
    stat(s, 7800000, Y0 + 1380000, 3720000, 1250000, "1 patent", "narrowed the claim to ring-VCO + closed PLL + open PDK + statistics", RED, size=24)
    stat(s, 7800000, Y0 + 2760000, 3720000, 1250000, "Open PDK", "everything reproducible: ngspice + sky130B, no licensed tools yet", TEAL, size=24)
    takeaway(s, "our contribution is the validated closed-loop study, not the idea of 2nd-order compensation.")
    add(s)

    # 6 method flow
    s = new_slide(P, "Methodology: Review 1 -> Review 2 -> Review 3", logo)
    picture(s, DIA / "method_flow.png", X0, Y0, W, 4300000)
    takeaway(s, "characterise, compensate, close the loop, then validate statistically.")
    add(s)

    # 7 architecture
    s = new_slide(P, "System Architecture", logo)
    picture(s, DIA / "pll_block_diagram.png", X0, Y0, W, 3000000)
    for i, (b, sm, c) in enumerate([("20 uA", "charge pump (behavioural), 5x until first divider edge", NAVY),
                                    ("3.6 kOhm", "R2 (retuned from 1.7 k); C1 68 pF, C2 6.8 pF", TEAL),
                                    ("5 stages", "current-starved ring VCO, tunes 18 -> 358 MHz", NAVY),
                                    ("/ 8", "divider (3 T flip-flops)", ORANGE)]):
        stat(s, X0 + i * 2760000, 4450000, 2600000, 1150000, b, sm, c, size=22)
    takeaway(s, "a standard integer-N loop, so every difference we see comes from the VCO and its compensation.")
    add(s)

    # 8 VCO
    s = new_slide(P, "VCO: Schematic & Tuning Curve", logo)
    picture(s, DIA / "vco_schematic.png", X0, Y0, 4300000, 3500000)
    ch = xy(s, 5000000, Y0 - 100000, 6500000, 3900000,
            [(c, D["tuning"][c], CC[c], "both") for c in ORDER],
            xt="Control voltage V_ctrl (V)", yt="VCO frequency (MHz)", xmin=0.6, xmax=1.4, ymin=0)
    legend(ch)
    stat(s, X0, 4950000, 2500000, 800000, "~890 MHz/V", "local K_VCO at 0.9 V (tt)", size=20)
    stat(s, X0 + 2650000, 4950000, 2500000, 800000, "18 -> 358 MHz", "over 0.65 -> 1.40 V (tt)", TEAL, size=20)
    takeaway(s, "V_ctrl sets the starving current; 0.9 V gives the 226 MHz target. K_VCO is ~890 MHz/V (an earlier 669 was wrong).", y=5850000)
    add(s)

    # 9 uncompensated drift
    s = new_slide(P, "Uncompensated Drift: PDK Corners", logo)
    ch = xy(s, X0, Y0 - 100000, 7200000, 4400000,
            [(f"{c} ({D['tempco_ppm'][c]:.0f} ppm/C)", pts(D["temp_dev"][c]), CC[c], "both") for c in ORDER]
            + [H.hline(0.3, -40, 85, "+/-0.3 % spec", GREEN), H.hline(-0.3, -40, 85, " ", GREEN)],
            xt="Temperature (C)", yt="Delta f vs 27 C (%)", xmin=-40, xmax=85)
    legend(ch, size=9)
    H.hide_legend_entry(ch, len(ORDER) + 1)
    stat(s, 8100000, Y0, 3400000, 1100000, "12.5x", "tempco spread, sf 112 to fs 1399 ppm/C", RED, size=24)
    stat(s, 8100000, Y0 + 1250000, 3400000, 1100000, "8.51 %", "tt frequency span, -40 to 85 C", NAVY, size=24)
    stat(s, 8100000, Y0 + 2500000, 3400000, 1100000, "112 - 1399", "ppm/C across the 5 PDK corners", TEAL, size=24)
    takeaway(s, "every corner is far outside +/-0.3 % without compensation, and the corners disagree by a factor of 12.")
    add(s)

    # 10-13 corner audit
    s = new_slide(P, "Why the Corner Numbers Looked Out of Range", logo)
    text(s, X0, Y0, W, 4400000, [
        ("In review, ff / tt / fs / sf values were not in the ranges they are meant to be. We audited it (ngspice 47, same PDK).", False, DARK),
        hd("Verdict: the simulations were fine. Three things were wrong in how we expected and described the corners."),
        hd("1. Labels"),
        "In this PDK sf = fast-NMOS / slow-PMOS and fs = slow-NMOS / fast-PMOS (measured). Earlier labels were reversed.",
        hd("2. Expectation"),
        "'Skew corners lie between ss and ff' is false in sky130: its skew corners shift each device further than ss/ff do. At the VCO's low bias this is amplified to about 9 %.",
        hd("3. Tempco is not bounded by ff and ss"),
        "Even in a clean textbook NMOS x PMOS set the tempco extremes sit at the skew corners."], size=15)
    takeaway(s, "nothing was broken; our expectation of the corner envelope was wrong.")
    add(s)

    s = new_slide(P, "Audit 1: PDK Skew Corners Are Wider Than ss / ff", logo)
    rows = [["Threshold shift vs tt (mV)", "ss", "ff", "PDK sf", "PDK fs"],
            ["NMOS Vth", "+25", "-25", "-36", "+36"],
            ["PMOS |Vth|", "+15", "-23", "+32", "-43"]]
    table(s, X0, Y0, 6200000, rows, [3, 1, 1, 1, 1], size=14, rowh=420000)
    text(s, X0, Y0 + 1500000, 6200000, 2500000, [
        "The extracted device thresholds of the skew corners move further than ss/ff.",
        "At the VCO bias (0.9 V, overdrive only ~0.2 V) an extra ~11 mV of Vth is ~9 % of drive current.",
        "At 1.8 V the same corners agree within 1 %."], size=13)
    stat(s, 7300000, Y0, 4220000, 1250000, "~9 %", "skew corners fall outside the ss-ff envelope at 0.9 V", RED, size=28)
    stat(s, 7300000, Y0 + 1400000, 4220000, 1250000, "< 1 %", "outside the envelope at 1.8 V", GREEN, size=28)
    stat(s, 7300000, Y0 + 2800000, 4220000, 1050000, "11 corners", "9 textbook + PDK sf + PDK fs", NAVY, size=24)
    takeaway(s, "the PDK defines its skew corners more aggressively than ss/ff, so they can fall outside them.")
    add(s)

    add(fig_slide(P, logo, "Audit 2: Device Current vs the ss-ff Envelope", FIG / "audit_device_envelope.png",
                  [("25.3 - 51.3 uA", "NMOS Id at 0.9 V: fs 25.3, ss 27.9, tt 36.9, ff 47.7, sf 51.3", NAVY),
                   ("397 - 480 uA", "NMOS Id at 1.8 V: skews within 1 % of ss/ff", TEAL),
                   ("8 - 9 %", "skews outside the ss-ff envelope at 0.9 V", RED)],
                  "at the low VCO bias, skew corners step outside ss/ff; at full gate drive they do not."))

    add(fig_slide(P, logo, "Audit 3: A Textbook 3 x 3 Corner Set", FIG / "audit_factorial_maps.png",
                  [("8 / 8", "ordering checks pass (f rises with N and P strength at 0.9 and 1.4 V; tempco falls with N strength)", GREEN),
                   ("178.0 - 279.2 MHz", "f at 0.9 V across the 9 textbook corners", NAVY),
                   ("306 - 1151", "textbook tempco range ppm/C (3.8x) vs PDK 5-corner 112 - 1399 (12.5x)", ORANGE)],
                  "with N and P varied independently the ordering is textbook-clean; only the PDK skews break the tempco range."))

    add(fig_slide(P, logo, "Audit 4: Tempco Follows NMOS Drive", FIG / "audit_tempco_vs_nmos.png",
                  [("-0.78 / -1.1", "Vth tempco mV/C, NMOS / PMOS (textbook -0.5 to -3)", NAVY),
                   ("-30 % / +14 %", "NMOS Id over 125 C at 1.8 V / 0.9 V", TEAL),
                   ("-816 ppm/C", "tt tempco at V_ctrl = 1.4 V (sign flips from +681 at 0.9 V)", RED)],
                  "weaker NMOS drive means larger positive tempco; the bias sits below the zero-tempco point."))

    s = new_slide(P, "Tempco Mechanism and Physics Checks", logo)
    text(s, X0, Y0, 6200000, 4400000, [
        hd("Why the frequency rises with temperature at 0.9 V"),
        "Vth falls with temperature (about -0.8 mV/C NMOS). With only ~0.2 V overdrive, that gain wins over the mobility loss, so drive current rises.",
        hd("Which device limits the ring depends on V_ctrl"),
        "At 0.9 V the NMOS dominates (36 % spread slow to fast N vs 10 % across P). At 1.4 V the PMOS dominates (26 % vs 9 %).",
        hd("Other checks that pass"),
        "Supply current is proportional to f in every corner (C ~ 27-30 fF per node). Closed-loop lock voltage equals the open-loop value for 226.2 MHz within 0.2 mV at all 11 corners."],
        size=13)
    for i, (b, sm, c) in enumerate([("36 % vs 10 %", "0.9 V: NMOS vs PMOS spread", NAVY),
                                    ("26 % vs 9 %", "1.4 V: PMOS vs NMOS spread", TEAL),
                                    ("0.2 mV", "lock V vs open-loop V, 11 corners", GREEN)]):
        stat(s, 7100000, Y0 + i * 1280000, 4420000, 1150000, b, sm, c, size=26)
    takeaway(s, "the tempco mechanism is overdrive, not N/P balance; every physics check agrees with theory.")
    add(s)

    # compensation
    s = new_slide(P, "Compensation Architecture", logo)
    picture(s, DIA / "compensation_arch.png", X0, Y0, W, 3000000)
    for i, (b, sm, c) in enumerate([("dV(T)", "correction added to V_ctrl before the VCO bias generator", NAVY),
                                    ("1st order", "k1 (T - 27): PTAT/CTAT equivalent", TEAL),
                                    ("2nd order", "adds k2 (T - 27)^2 for the curvature", NAVY),
                                    ("Ideal", "behavioural source; transistor-level network is future work", ORANGE)]):
        stat(s, X0 + i * 2760000, 4450000, 2600000, 1150000, b, sm, c, size=22)
    takeaway(s, "the correction is a voltage added to V_ctrl: v_eff = v + k1 (T-27) + k2 (T-27)^2.")
    add(s)

    s = new_slide(P, "Compensation Method: Fit, Then Verify", logo)
    text(s, X0, Y0, 6600000, 4400000, [
        hd("scripts/comp_factorial.py, applied per corner"),
        "1. Measure f(T, V) at V = 0.85 / 0.9 / 0.95 V for 6 temperatures.",
        "2. Solve the exact voltage shift dV*(T) that restores f(27 C).",
        "3. Least-squares fit k1 (1st order) and k1, k2 (2nd order).",
        "4. Verify by simulating the compensated VCO at the 6 fit temperatures and at 4 held-out temperatures (-30, 10, 40, 70 C).",
        hd("Why Review 2's numbers changed"),
        "Review 2 used a wrong K_VCO (669 instead of ~890 MHz/V) and an analytic first-order guess. The measured fit replaces it: tt span 1.12 % -> 0.20 %."], size=13)
    stat(s, 7500000, Y0, 4020000, 1150000, "-1.543e-4", "k1 at tt (V/C)", NAVY, size=24)
    stat(s, 7500000, Y0 + 1280000, 4020000, 1150000, "4.18e-7", "k2 at tt (V/C^2)", TEAL, size=24)
    stat(s, 7500000, Y0 + 2560000, 4020000, 1150000, "4 held-out", "temperatures never used in the fit", ORANGE, size=24)
    takeaway(s, "coefficients come from measured data and are checked at temperatures the fit never saw.")
    add(s)

    r = CRN["ntpt"]
    add(fig_slide(P, logo, "Compensation at the Typical (tt) Corner", FIG / "final_comp_tt.png",
                  [("8.51 -> 0.20 %", "span: baseline -> 2nd order (1st order 0.67 %)", NAVY),
                   ("0.16 %", "max deviation; +/-0.3 % spec MET", GREEN),
                   ("0.09 %", "max deviation at held-out temperatures", TEAL)],
                  "at tt the compensated VCO stays within +/-0.16 % over -40 to 85 C."))

    add(fig_slide(P, logo, "Compensation on 11 Corners", FIG / "final_comp_corners.png",
                  [("8 / 11", "corners meet +/-0.3 % (7 of 9 textbook)", GREEN),
                   ("31 - 44x", "compensation ratio (PDK sf 15.9x)", NAVY),
                   ("0.03 - 0.25 %", "held-out max deviation", TEAL)],
                  "8 of 11 corners meet the spec; the slow-NMOS corners are the hardest."))

    s = new_slide(P, "Compensation Results: All 11 Corners", logo)
    rows = [["Corner", "f 0.9 V (MHz)", "Base span %", "1st %", "2nd %", "Max dev %", "Held-out %", "Ratio C", "+/-0.3 %"]]
    bad = []
    for i, k in enumerate(KEYS, 1):
        c = CRN[k]
        rows.append([LABEL[k], c["f(0.9 V, 27 C) MHz"], c["baseline span %"], c["1st-order span %"], c["2nd-order span %"],
                     c["2nd max dev %"], c["held-out max dev %"], c["C = base/2nd"], "yes" if c["meets +/-0.3 %"] == "yes" else "NO"])
        if c["meets +/-0.3 %"] != "yes":
            bad.append(i)
    table(s, X0, Y0 - 50000, W, rows, [3.2, 1.3, 1.2, 0.9, 0.9, 1.1, 1.2, 1, 1], size=11, rowh=335000, bad=bad)
    takeaway(s, "NOT met: N slow / P typ 0.301 (borderline), N slow / P fast 0.325, PDK fs 0.360. Nominal, ideal source.", y=5750000)
    add(s)

    # Monte Carlo
    s = new_slide(P, "Monte Carlo (n = 72 chips, tt): Mean |Drift|", logo)
    ch = bar(s, X0, Y0 - 50000, 6300000, 4300000, ["Baseline 20.11", "1st order 3.25", "2nd order 2.76"],
             [("mean |f(85C) - f(-40C)| (MHz)", [20.11, 3.25, 2.76], [GREY, TEAL, NAVY])],
             xt=None, yt="Mean |drift| (MHz), 95 % CI bars", ymin=0)
    ch.has_legend = False
    error_bars(ch.series[0], [(21.04 - 19.18) / 2, (3.87 - 2.63) / 2, (3.19 - 2.33) / 2])
    stat(s, 7300000, Y0, 4220000, 1100000, "-17.35 MHz", "baseline -> 2nd order, CI [16.37, 18.34], 72/72 chips", GREEN, size=24)
    stat(s, 7300000, Y0 + 1250000, 4220000, 1100000, "-16.86 MHz", "baseline -> 1st order, CI [15.57, 18.16], 71/72 chips", TEAL, size=24)
    stat(s, 7300000, Y0 + 2500000, 4220000, 1250000, "0.49 MHz", "1st vs 2nd order: CI [-0.26, 1.24], NOT significant (2nd smaller in 40/72)", RED, size=24)
    takeaway(s, "compensation cuts drift ~7x (72/72 chips), but the 2nd-order term adds no detectable benefit under mismatch.")
    add(s)

    add(fig_slide(P, logo, "Monte Carlo: The Mismatch Floor", FIG / "final_mc.png",
                  [("~3 MHz", "residual drift (~1.2-1.4 % of 226 MHz)", RED),
                   ("0.2 - 0.7 %", "nominal residual without mismatch", GREEN),
                   ("Per-chip trim", "needed to reach +/-0.3 % across chips (future work)", ORANGE)],
                  "transistor mismatch makes each chip drift differently, and one global correction cannot remove it."))

    # closed loop
    s = new_slide(P, "Loop Dynamics: Found, Fixed, Verified", logo)
    text(s, X0, Y0, 6000000, 4400000, [
        hd("Found"),
        "The Review-1 loop design dropped a 1/(2 pi) factor in omega_n^2, so damping was zeta ~0.33, not the 0.72 claimed. The original loop had not settled at ss/ff/sf/fs within 2 us.",
        hd("Fixed"),
        "R2 retuned 1.7 k -> 3.6 kOhm: zeta ~0.70, phase margin ~65-70 deg (100 x zeta rule, not a Bode plot).",
        hd("Cost"),
        "Cold start is slower (tt 839 -> 1351 ns) but in spec; coarse start faster (577 -> 503 ns)."], size=13)
    for i, (b, sm, c) in enumerate([("zeta 0.33 -> 0.70", "damping", NAVY), ("8.6 -> 0.35 mV", "V_ctrl ripple at lock", TEAL),
                                    ("0.78 -> 0.008 %", "final frequency error", GREEN)]):
        stat(s, 6900000, Y0 + i * 1280000, 4620000, 1150000, b, sm, c, size=24)
    takeaway(s, "we caught a design-equation error ourselves; the retuned loop is well damped and quiet.")
    add(s)

    s = new_slide(P, "Cold-Start Fix: Boosted Charge Pump", logo)
    picture(s, DIA / "acq_boost.png", X0, Y0, W, 2400000)
    text(s, X0, 3850000, 6300000, 1900000, [
        hd("Problem"),
        "From 0 V the VCO is in deep subthreshold (0.6 MHz at 0.5 V, 6.5 MHz at 0.6 V vs 226 MHz target), so the divider clock is effectively absent and the pump-up is slow.",
        hd("Fix"),
        "Charge pump runs at 5 x I_cp until the divider's first edge, then drops to 20 uA."], size=12)
    stat(s, 7200000, 3900000, 2150000, 1250000, "1351 ns", "tt lock from 0 V", NAVY, size=24)
    stat(s, 9500000, 3900000, 2020000, 1250000, "503 ns", "tt lock from 0.85 V", TEAL, size=24)
    takeaway(s, "the boosted pump removes the dead zone at start-up; lock is inside 2 us.")
    add(s)

    add(fig_slide(P, logo, "Cold-Start Traces", FIG / "final_cold_traces.png",
                  [("1283 - 1675 ns", "cold-start lock, 11 corners, 27 C", NAVY),
                   ("<= 0.8 mV", "V_ctrl ripple in every run", GREEN),
                   ("226.17 MHz", "= 8 x f_ref (-0.01 %) in every run", TEAL)],
                  "from 0 V the control voltage climbs and settles to its corner-specific value."))

    s = new_slide(P, "Lock Time Across 11 Corners (27 C)", logo)
    cats = [SHORT[k] for k in KEYS]
    bar(s, X0, Y0 - 50000, 7700000, 4350000, cats,
        [("Coarse start (0.85 V)", [float(LOOP[f"{k}_T27_coarse"]["lock ns"]) for k in KEYS], NAVY),
         ("Cold start (0 V, boosted CP)", [float(LOOP[f"{k}_T27_cold"]["lock ns"]) for k in KEYS], ORANGE)],
        xt=None, yt="Lock time (ns), spec < 2000", ymin=0)
    legend(s.shapes[-1].chart)
    stat(s, 8600000, Y0, 2920000, 1100000, "177 - 967 ns", "coarse start", NAVY, size=22)
    stat(s, 8600000, Y0 + 1250000, 2920000, 1100000, "1283 - 1675 ns", "cold start", ORANGE, size=22)
    stat(s, 8600000, Y0 + 2500000, 2920000, 1100000, "22 / 22", "runs inside 2 us at 27 C", GREEN, size=22)
    takeaway(s, "all 11 corners lock inside 2 us at 27 C; slow-NMOS corners take longest. (Ns = N slow, Pf = P fast, etc.)")
    add(s)

    add(fig_slide(P, logo, "Does the Loop Lock at the Right Voltage?", FIG / "final_vlock_check.png",
                  [("< 0.2 mV", "lock V vs open-loop V for 226.2 MHz, 11 corners", GREEN),
                   ("0.849 - 0.968 V", "lock V_ctrl range, ff to ss at 27 C", NAVY),
                   ("+11.9 / -7.5 mV", "tt lock-V shift at -40 / 85 C = fitted compensation (+12.2 / -7.5)", TEAL)],
                  "the PLL absorbs the VCO tempco; at tt this independently confirms the compensation fit."))

    add(fig_slide(P, logo, "Temperature: Cold-Start Lock at -40, 27, 85 C", FIG / "final_lock_temp.png",
                  [("2083 ns", "ss at 85 C: MISSES the 2 us spec by 4 %", RED),
                   ("1319 / 1418 ns", "tt at -40 / 85 C", GREEN),
                   ("1249 / 1317 ns", "ff at -40 / 85 C", GREEN)],
                  "worst case ss / 85 C misses the lock spec slightly; every other case is inside it."))

    # benchmark
    s = new_slide(P, "Benchmark vs Literature (Honest)", logo)
    rows = [["Design", "Process / basis", "Residual drift", "Lock time"],
            ["This work, 2nd order (tt, nominal)", "sky130, simulation, ideal source", "+/-0.16 % max dev (0.20 % span), -40 to 85 C", "1351 ns (0 V cold start)"],
            ["This work, 2nd order (MC n = 72)", "sky130, simulation, mismatch on", "~2.76 MHz = ~1.2 % |drift|, 85 vs -40 C", "n/a"],
            ["Souliotis 2022 [1]", "180 nm, measured silicon", "-1.5 % / +0.4 %, -40 to 120 C", "not reported"],
            ["Stancu 2023 [2]", "simulation, 50 MHz clock", "1.6 % temperature error", "n/a"],
            ["Geng 2018 [3]", "180 nm, measured silicon", "not reported as drift", "< 2 us"]]
    table(s, X0, Y0, W, rows, [3, 2.6, 3.4, 2], size=12, rowh=520000, left=2)
    text(s, X0, 4600000, W, 1200000, [
        "Nominal simulation of an ideal network is competitive, but it is not like-for-like with measured silicon or with a transistor-level circuit. With chip-to-chip mismatch our residual (~1.2 %) is the same order as the published numbers.",
        "No jitter or phase-noise number yet (no PSS/Pnoise in open tools), so [3]'s -117 dBc/Hz cannot be compared."], size=12)
    takeaway(s, "in the right ballpark, not proof of parity with fabricated designs.", y=5850000)
    add(s)

    # corrections
    s = new_slide(P, "What We Found and Corrected", logo)
    rows = [["#", "Earlier claim", "Corrected"],
            ["1", "sf / fs labels (sf = slow N)", "reversed: PDK sf = fast N / slow P (measured)"],
            ["2", "Tempco set by N/P balance", "set by NMOS overdrive; sign flips at 1.4 V"],
            ["3", "K_VCO = 669 MHz/V", "~890 MHz/V (887 local slope at 0.9 V)"],
            ["4", "Hard oscillation floor at 0.63 V", "none: 0.63 MHz at 0.5 V, 6.5 MHz at 0.6 V (window artefact)"],
            ["5", "Loop zeta 0.72 (1/(2 pi) dropped)", "zeta 0.33 -> retuned R2 = 3.6 k gives ~0.70"],
            ["6", "Compensation coefficients (wrong K_VCO)", "refit: tt 0.16 % max dev, 8 / 11 corners meet spec"],
            ["7", "2nd order significantly better (n = 30)", "n = 72, corrected fit: not significant; mismatch floor ~1.2-1.4 %"],
            ["8", "PDK skew corners lie inside ss-ff", "wider than ss/ff; textbook 3 x 3 set added"],
            ["9", "Closed loop fine in all cases", "ss / 85 C cold start 2083 ns misses 2 us by 4 %"]]
    table(s, X0, Y0 - 50000, W, rows, [0.4, 4, 6], size=12, rowh=400000, left=3)
    takeaway(s, "each fix came from our own audit and each is traceable to a re-run.", y=5800000)
    add(s)

    # limitations
    txt_slide(P, logo, "Known Limitations", [
        "Compensation network and charge pump are behavioural (ideal) models; no transistor-level versions yet.",
        "No layout, DRC, LVS or post-layout simulation yet.",
        "No jitter or phase-noise number (no PSS/Pnoise in open tools).",
        "Phase margin is the 100 x zeta rule, not a Bode plot.",
        "Only R2 = 1.7 / 3.0 / 3.6 kOhm simulated at transistor level; 3.6 k is not claimed optimal.",
        "Cold start simulated on 11 corners at 27 C and at 3 corners for +/-temperature, not all 11 x all T.",
        "Compensation is fitted nominally at V_ctrl = 0.9 V; each corner actually locks at 0.85 - 0.97 V.",
        "Mismatch limits compensated drift to ~1.2-1.4 %; the H1 test and Spectre arm are pending."], size=17,
        take="we state plainly what is ideal, nominal, or still missing.")
    S.append(P.slides[-1])

    # next phase
    txt_slide(P, logo, "Remaining Work / Next Phase", [
        hd("Circuit"),
        "Transistor-level charge pump and compensation network (PTAT/CTAT + squarer), then re-run the loop.",
        "Define the compensation at each corner's own lock voltage; per-chip trim to beat the mismatch floor.",
        hd("Physical design"),
        "Layout in Magic, DRC, LVS (netgen, built), post-layout simulation.",
        hd("Validation and release"),
        "Virtuoso/Spectre cross-check arm, PSS/Pnoise, H1 regression (needs department Virtuoso access - to confirm with Dr. Bindu).",
        "Spec sign-offs (V_ctrl floor revision, +/-0.15 ppm restatement); paper and artifact release."], size=16,
        take="the behavioural result is solid; the next phase makes it real hardware-shaped.")
    S.append(P.slides[-1])

    # at a glance
    s = new_slide(P, "Final Review at a Glance", logo)
    items = [("8 / 11", "corners meet +/-0.3 % (tt 0.16 % max dev)", GREEN),
             ("8.51 -> 0.20 %", "tt span, baseline -> 2nd order", NAVY),
             ("31 - 44x", "compensation ratio (PDK sf 15.9x)", TEAL),
             ("20.1 -> 2.8 MHz", "Monte Carlo |drift|, n = 72", NAVY),
             ("1.2 - 1.4 %", "mismatch floor (per-chip trim needed)", RED),
             ("177 - 967 ns", "coarse-start lock, 11 corners, 27 C", TEAL),
             ("1283 - 1675 ns", "cold-start lock, 11 corners, 27 C", ORANGE),
             ("2083 ns", "ss / 85 C cold start: misses 2 us by 4 %", RED)]
    for i, (b, sm, c) in enumerate(items):
        stat(s, X0 + (i % 4) * 2760000, Y0 + (i // 4) * 2000000, 2600000, 1800000, b, sm, c, size=22)
    text(s, X0, 5400000, W, 400000, ["zeta ~0.70 (R2 = 3.6 k); 15 references; 1 patent narrowed the claim; 9 corrections found by our own audit."], size=13)
    takeaway(s, "a verified, honest closed-loop study with its limits stated.", y=5850000)
    S.append(s)
    return S


def main(base=BASE, out=OUT):
    prs = Presentation(str(base))
    logo = next(sh for sh in prs.slides[1].shapes if sh.shape_type == 13).image.blob
    old = list(prs.slides)
    # title slide text
    for sh in old[0].shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if "Review-II" in r.text:
                        r.text = "Final Review (Review-III)  ·  28.10.2026"
    new = build(prs, logo)
    # order: title, new[0..], lit slides after problem
    lit1, lit2, refs, thanks = old[2], old[3], old[25], old[26]
    # new indices: 0 problem, 1 novelty, 2 method, 3 arch, ... last = glance
    novelty = new[1]
    body = new[2:]
    order = [old[0], new[0], lit1, lit2, novelty] + body + [refs, thanks]
    lst = prs.slides._sldIdLst
    by = {sl.slide_id: el for sl, el in zip(prs.slides, list(lst))}
    keep = {sl.slide_id for sl in order}
    for sid, el in by.items():
        if sid not in keep:
            prs.part.drop_rel(el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
            lst.remove(el)
    for sl in order:
        lst.append(by[sl.slide_id])
    for n, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame and sh.left == 11460175 and sh.top == 6446520:
                sh.text_frame.paragraphs[0].runs[0].text = str(n)
    prs.save(str(out))
    print("wrote", out, len(prs.slides), "slides")


if __name__ == "__main__":
    main(*(Path(a) for a in sys.argv[1:3]))
